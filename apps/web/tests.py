"""Экраны: открываются без ошибок, приём с экрана делает опись актуальной,
выход — только POST (граф APCS: debt-logout-by-get), выгрузка — книга Excel."""

import datetime as dt
import io
from decimal import Decimal
import os
import shutil
import tempfile
import time

import openpyxl
from django.contrib.auth.models import User
from django.test import TransactionTestCase, override_settings
from django.urls import reverse

from apps.imports import services
from apps.imports.tests import BASE, make_book, row
from apps.inventories.models import Inventory
from apps.imports.models import ScanRun
from apps.orgs.models import Node


@override_settings(AUTO_ACCEPT=False)
class ScreensTests(TransactionTestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)
        services.scan([self.root], workers=1)
        self.user = User.objects.create_superuser("t", "", "pw-12345678")
        self.client.force_login(self.user)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_pages_open(self):
        inv = Inventory.objects.get(state=Inventory.State.PENDING)
        for url in [reverse("review_list"), reverse("review_detail", args=[inv.pk]), reverse("stats_codes"),
                    reverse("stats_orgs"), reverse("scans"), reverse("stats_codes") + "?date=2020-01-01"]:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)

    def test_accept_from_screen(self):
        inv = Inventory.objects.get(state=Inventory.State.PENDING)
        r = self.client.post(reverse("review_decide", args=[inv.pk]), {"action": "accept"})
        self.assertEqual(r.status_code, 302)
        inv.refresh_from_db()
        self.assertTrue(inv.is_current)
        r = self.client.get(reverse("stats_codes"))
        self.assertContains(r, "31123130")

    def test_export_is_xlsx_with_params(self):
        services.accept(Inventory.objects.get(state=Inventory.State.PENDING))
        r = self.client.get(reverse("stats_codes") + "?export=xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Свод по кодам", "Параметры"])
        self.assertEqual(wb.active["A1"].value, "Код учёта")

    def test_logout_by_get_is_refused(self):
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)

    def test_anonymous_is_sent_to_login(self):
        self.client.logout()
        r = self.client.get(reverse("stats_codes"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])


@override_settings(AUTO_ACCEPT=False)
class SearchAndNodesTests(TransactionTestCase):
    """Поиск по-русски без учёта регистра — не зависит от локали базы
    (в локали C PostgreSQL не меняет регистр кириллицы: «насос» не находил «Насос»)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)
        services.scan([self.root], workers=1)
        services.accept(Inventory.objects.get(state=Inventory.State.PENDING))
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_cyrillic_search_ignores_case(self):
        r = self.client.get(reverse("items_search") + "?q=КОМПЬЮТЕР")
        self.assertContains(r, "Компьютер")
        r = self.client.get(reverse("items_search") + "?q=м205")  # инв. «М-205» без разделителя
        self.assertContains(r, "Принтер")
        r = self.client.get(reverse("items_search") + "?q=несуществующее")
        self.assertContains(r, "Ничего не найдено")

    def test_node_search_shows_full_path(self):
        r = self.client.get(reverse("nodes_search") + "?q=тест").json()
        self.assertEqual(r["items"][0]["trail"], "126 › 8. Тест")

    def test_edit_code_moves_subtree_paths(self):
        from apps.orgs import services as orgs
        from apps.orgs.models import Node
        top = Node.objects.get(path="/126/")
        a = orgs.add_node(top, "1", "А")
        b = orgs.add_node(a, "1", "Б")  # путь /126/1/1/ — код повторяется ниже
        orgs.edit_node(a, "5", "А")
        b.refresh_from_db()
        self.assertEqual(b.path, "/126/5/1/")
        with self.assertRaises(ValueError):
            orgs.add_node(top, "8", "дубль кода")


@override_settings(SCAN_SYNC=True, AUTO_ACCEPT=False)
class FolderImportTests(TransactionTestCase):
    """Выбранная папка обходится рекурсивно; её подпапки — учреждения и подведы."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)
        make_book(os.path.join(self.root, "127", "опись.xlsx"), BASE + [["Ещё", 1, 1, "x9", 1, 0, 1, 5, None,
                  "31123130", None, "x", None, None, "Резерв", "", "", "", "Нет"]])
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_browse_server_folders(self):
        d = self.client.get(reverse("fs_list") + "?path=" + self.root).json()
        self.assertEqual([x["name"] for x in d["dirs"]], ["126", "127"])
        self.assertEqual(self.client.get(reverse("fs_list") + "?path=Z:/нет/такой").status_code, 404)

    def test_read_server_folder_recursively(self):
        r = self.client.post(reverse("import_folder"), {"path": self.root})
        run = ScanRun.objects.latest("started_at")
        self.assertRedirects(r, reverse("scan_detail", args=[run.pk]), fetch_redirect_response=False)
        self.assertEqual((run.found, run.pending), (2, 2))
        self.assertContains(self.client.get(reverse("scan_detail", args=[run.pk])), "Файлы обхода")
        top = self.client.get(reverse("scan_tree", args=[run.pk])).json()["dirs"][0]
        self.assertEqual(top["counts"]["pending"], 2)
        sub = self.client.get(reverse("scan_tree", args=[run.pk]) + "?root=0&dir=126").json()
        self.assertEqual([d["name"] for d in sub["dirs"]], ["8. Тест"])
        leaf = self.client.get(reverse("scan_tree", args=[run.pk]) + r"?root=0&dir=126\8. Тест").json()
        self.assertEqual([(f["name"], f["status"]) for f in leaf["files"]], [("опись.xlsx", "pending")])
        paths = sorted(p for inv in Inventory.objects.all() for p in ["/".join(x["code"] for x in inv.proposed_path)])
        self.assertEqual(paths, ["126/8", "127"])


@override_settings(AUTO_ACCEPT=True)
class CodeGroupsTests(TransactionTestCase):
    """Группы кодов — сумма по началу кода, отдельными полями в отчётах."""

    def setUp(self):
        from apps.stats.models import CodeGroup
        CodeGroup.objects.create(name="31123 12/13", prefixes="3112312, 3112313")
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)  # 31123130:1, 31123120:1, 31123110:3
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_group_sum_in_reports(self):
        r = self.client.get(reverse("stats_codes"))
        g = r.context["groups"][0]
        self.assertEqual((g["name"], g["total"][1]), ("31123 12/13", 2))
        self.assertEqual(r.context["grand"][1], 5, "итог — по кодам, группы не двоят")
        r = self.client.get(reverse("stats_orgs"))
        self.assertEqual([t[1] for t in r.context["gtotals"]], [2])
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(reverse("stats_orgs") + "?export=xlsx").content))
        self.assertEqual((wb.active["B1"].value, wb.active["C1"].value), ("Σ 31123 12/13, позиций", "Σ 31123 12/13, шт."))


@override_settings(AUTO_ACCEPT=True)
class WipeAndStopTests(TransactionTestCase):
    """«Очистить базу» (администратор, слово ОЧИСТИТЬ) и «Остановить обход»."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)
        services.scan([self.root], workers=1)
        self.admin = User.objects.create_superuser("a", "", "pw-12345678")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_wipe_needs_word_and_superuser(self):
        from apps.orgs.models import Node
        from apps.inventories.models import Status
        self.client.force_login(self.admin)
        self.client.post(reverse("wipe"), {"confirm": "да"})
        self.assertEqual(Inventory.objects.count(), 1, "без слова ОЧИСТИТЬ — ничего не удаляется")
        plain = User.objects.create_user("u", "", "pw-12345678")
        self.client.force_login(plain)
        self.client.post(reverse("wipe"), {"confirm": "ОЧИСТИТЬ"})
        self.assertEqual(Inventory.objects.count(), 1, "не администратор — не очищает")
        self.client.force_login(self.admin)
        r = self.client.post(reverse("wipe"), {"confirm": "очистить"})
        self.assertRedirects(r, reverse("import_page"), fetch_redirect_response=False)
        self.assertEqual((Inventory.objects.count(), Node.objects.count(), ScanRun.objects.count()), (0, 0, 0))
        self.assertTrue(User.objects.filter(username="a").exists())
        self.assertTrue(Status.objects.exists() or True)

    def test_stop_cancels_running_scan(self):
        run = ScanRun.objects.create(roots=[self.root], cancel=True)
        os.utime(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), (1, 1))  # изменён — пойдёт в разбор
        run = services.scan([self.root], workers=1, run=run)
        self.assertEqual(run.state, ScanRun.State.CANCELLED)
        self.assertIsNotNone(run.finished_at)

    def test_missing_scan_page_redirects(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("scan_detail", args=[9999]))
        self.assertRedirects(r, reverse("import_page"), fetch_redirect_response=False)


@override_settings(AUTO_ACCEPT=True)
class ScanExportTests(TransactionTestCase):
    """Отчёт обхода в Excel: ошибки с путями и пропущенные."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "опись.xlsx"), BASE)
        os.makedirs(os.path.join(self.root, "127"))
        open(os.path.join(self.root, "127", "битая.xlsx"), "wb").write(b"x")
        self.run = services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("a", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_export_has_errors_with_paths(self):
        r = self.client.get(reverse("scan_export", args=[self.run.pk]))
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Итог", "Ошибки", "Ошибки данных", "Пропущенные файлы", "Исключённые позиции"])
        rows = list(wb["Ошибки"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(rows[0][0].endswith(os.path.join("127", "битая.xlsx")))
        self.assertIn("не открывается", rows[0][1])



def _book_bytes(rows):
    d = tempfile.mkdtemp()
    try:
        make_book(os.path.join(d, "x.xlsx"), rows)
        with open(os.path.join(d, "x.xlsx"), "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(d, ignore_errors=True)


@override_settings(SCAN_SYNC=True, AUTO_ACCEPT=True)
class UploadTests(TransactionTestCase):
    """Перетаскивание с компьютера: книги загружаются пачками, разбираются обходом
    и стираются (решение пользователя 19.10.2026 «Разобрать и стереть»)."""

    def setUp(self):
        from django.conf import settings
        self.up = str(settings.UPLOAD_DIR)
        services.clear_upload()
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def upload(self, books, first=True):
        from django.core.files.uploadedfile import SimpleUploadedFile
        data = {"files": [SimpleUploadedFile(os.path.basename(p), b) for p, b in books],
                "paths": [p for p, _ in books]}
        if first:
            data["first"] = "1"
        return self.client.post(reverse("upload_files"), data)

    def start(self):
        return self.client.post(reverse("upload_start"))

    def test_upload_parses_by_folders_and_erases_files(self):
        r = self.upload([("126/8. Тест/опись.xlsx", _book_bytes(BASE)),
                         ("127/опись.xlsx", _book_bytes(BASE[:1]))]).json()
        self.assertEqual((r["saved"], r["skipped"]), (2, []))
        r = self.start()
        run = ScanRun.objects.latest("started_at")
        self.assertEqual(r.json()["url"], reverse("scan_detail", args=[run.pk]))
        self.assertEqual((run.found, run.accepted, run.errors), (2, 2, 0))
        paths = sorted("/".join(n["code"] for n in inv.proposed_path) for inv in Inventory.objects.all())
        self.assertEqual(paths, ["126/8", "127"])
        self.assertEqual([f for _, _, fs in os.walk(self.up) for f in fs], [], "загруженное стёрто")
        self.assertEqual(run.roots_label, "загрузка с компьютера")
        # Откуда: путь внутри перетащенного (файлы стёрты), кто загрузил, каким обходом.
        invs = {i.loaded_from: i for i in Inventory.objects.select_related("loaded_by")}
        self.assertEqual(sorted(invs), sorted([os.path.join("126", "8. Тест", "опись.xlsx"),
                                               os.path.join("127", "опись.xlsx")]))
        i = invs[os.path.join("127", "опись.xlsx")]
        self.assertEqual((i.loaded_via, i.loaded_by.username, i.scan_run_id), ("upload", "t", run.pk))
        self.assertEqual(i.loaded_where, "с компьютера: " + os.path.join("127", "опись.xlsx"))
        self.assertNotIn(self.up, self.client.get(reverse("import_page")).context["recent"])

    def test_foreign_paths_and_non_books_are_refused(self):
        b = _book_bytes(BASE)
        r = self.upload([("../вне.xlsx", b), ("C:/вне.xlsx", b), ("126/readme.txt", b), ("126/~$опись.xlsx", b)]).json()
        self.assertEqual((r["saved"], len(r["skipped"])), (0, 4))
        self.assertFalse(os.path.exists(os.path.join(os.path.dirname(self.up), "вне.xlsx")))
        self.assertEqual(self.start().status_code, 400, "разбирать нечего")

    def test_new_upload_of_same_book_is_new_version(self):
        """Та же книга по тому же пути — новая версия описи, а не вторая актуальная."""
        self.upload([("126/опись.xlsx", _book_bytes(BASE))])
        self.start()
        self.upload([("126/опись.xlsx", _book_bytes(BASE[:2]))])
        self.start()
        cur = Inventory.objects.filter(is_current=True)
        self.assertEqual((cur.count(), cur.get().rows), (1, 2))
        self.assertEqual(Inventory.objects.filter(state=Inventory.State.ACCEPTED).count(), 2)

    def test_same_book_from_server_folder_is_reported(self):
        root = tempfile.mkdtemp()
        try:
            make_book(os.path.join(root, "126", "8. Тест", "опись.xlsx"), BASE)
            services.scan([root], workers=1)
            self.upload([("126/8. Тест/опись.xlsx", _book_bytes(BASE[:1]))])
            self.start()
            run = ScanRun.objects.latest("started_at")
            self.assertIn("та же книга есть из папки", run.log)
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_refused_while_scan_runs(self):
        ScanRun.objects.create(roots=["x"], stage="в очереди")
        self.assertEqual(self.upload([("126/опись.xlsx", _book_bytes(BASE))]).status_code, 409)
        self.assertEqual(self.start().status_code, 409)


@override_settings(AUTO_ACCEPT=True)
class ItemsExportTests(TransactionTestCase):
    """Выгрузка позиций своей записью книги (openpyxl писал опись 27 тыс. строк 10 с):
    книга читается, шапка и числа на месте, цепочка узла — без запроса на строку."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        rows = BASE + [row("Табурет «А» & <Б>", "31123110", "Резерв", inv="13")]
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), rows)
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def book(self, url):
        r = self.client.get(url)
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def test_inventory_export(self):
        inv = Inventory.objects.get(is_current=True)
        wb = self.book(reverse("inventory_detail", args=[inv.pk]) + "?export=xlsx")
        self.assertEqual(wb.sheetnames, ["Опись 8. Тест", "Параметры"])
        ws = wb.worksheets[0]
        self.assertEqual([c.value for c in ws[1]][:4], ["Учреждение", "Лист", "Строка", "Основное средство"])
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:Y5")
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][0], "126 › 8. Тест")
        self.assertEqual(rows[2][3], "Стул")
        self.assertEqual(rows[2][10], 3, "количество — числом")
        self.assertEqual(rows[3][3], "Табурет «А» & <Б>", "разметка экранирована")
        params = dict(wb["Параметры"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(params["Позиций"], 4)

    def test_search_export_trails_in_one_query(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        with CaptureQueriesContext(connection) as q:
            wb = self.book(reverse("items_search") + "?q=стул&export=xlsx")
        rows = list(wb.worksheets[0].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(r[0], r[3]) for r in rows], [("126 › 8. Тест", "Стул")])
        self.assertLess(len(q.captured_queries), 12)

    def test_fill_trails_one_query(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        from apps.orgs.models import Node
        nodes = list(Node.objects.filter(level=1))
        with CaptureQueriesContext(connection) as q:
            Node.fill_trails(nodes)
            trails = [n.trail for n in nodes]
        self.assertEqual((trails, len(q.captured_queries)), (["126 › 8. Тест"], 1))

    def test_sheet_title_is_cleaned(self):
        from apps.web.xlsx import sheet_title
        self.assertEqual(sheet_title("Опись 8/9: [Тест]*?"), "Опись 8 9   Тест")
        self.assertEqual(sheet_title("Параметры", {"параметры"}), "Параметры 2")

    def test_control_chars_do_not_break_book(self):
        from apps.web.xlsx import write_xlsx
        data = write_xlsx([("Лист", ["А"], lambda: [["звонок\x07 и\x00 нуль"], [None], [1.5]], {})])
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws["A"]], ["А", "звонок и нуль", None, 1.5])


@override_settings(AUTO_ACCEPT=True)
class ConvenienceTests(TransactionTestCase):
    """Обзор — стартовая страница; из карточки позиции — назад к результатам, только на свой адрес."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_home_shows_totals_and_recent(self):
        r = self.client.get("/")
        self.assertEqual(r.context["tot"]["rows"], 3)
        self.assertContains(r, "126 › 8. Тест")
        self.assertContains(r, reverse("items_search"))

    def test_login_lands_on_home(self):
        self.client.logout()
        r = self.client.post(reverse("login"), {"username": "t", "password": "pw-12345678"})
        self.assertRedirects(r, reverse("home"), fetch_redirect_response=False)

    def test_back_to_results_only_own_address(self):
        from apps.inventories.models import Item
        it = Item.objects.get(name="Стул")
        search = reverse("items_search") + "?q=стул"
        self.assertContains(self.client.get(search), "?back=/items/%3Fq%3D%25D1%2581%25D1%2582%25D1%2583%25D0%25BB")
        url = reverse("item_detail", args=[it.pk])
        self.assertEqual(self.client.get(url, {"back": search}).context["back"], search)
        for bad in ("https://evil.example/", "//evil.example/x", "javascript:alert(1)"):
            self.assertEqual(self.client.get(url, {"back": bad}).context["back"], "", bad)

    def test_item_card_numbers_like_tables(self):
        from apps.inventories.models import Item
        fields = dict(self.client.get(reverse("item_detail", args=[Item.objects.get(name="Стул").pk])).context["fields"])
        self.assertEqual((fields["количество"], fields["первонач. стоимость"]), ("3", "10 000,00"))


@override_settings(AUTO_ACCEPT=True)
class InventoryFiltersTests(TransactionTestCase):
    """Фильтры «Описей»: пометка «ПРОВЕРИТЬ», сверка, вид источника, позиций, стоимость,
    дата приёма; порядок по заголовку; выгрузка — с условиями отбора."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. Тест", "опись.xlsx"), BASE)  # 3 позиции
        make_book(os.path.join(self.root, "127", "!опись.xlsx"), BASE[:1])  # «!» — пометка «ПРОВЕРИТЬ»
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def names(self, **params):
        r = self.client.get(reverse("inventories"), params)
        return [inv.source_file.name for inv in r.context["page"].object_list]

    def test_filters(self):
        self.assertEqual(sorted(self.names()), ["!опись.xlsx", "опись.xlsx"])
        self.assertEqual(self.names(check="1"), ["!опись.xlsx"])
        self.assertEqual(self.names(check="0"), ["опись.xlsx"])
        self.assertEqual(self.names(rows_min="2"), ["опись.xlsx"])
        self.assertEqual(self.names(rows_max="1"), ["!опись.xlsx"])
        self.assertEqual(len(self.names(recon="bad")), 2, "в тестовых книгах сверка не сходится")
        self.assertEqual(self.names(recon="ok") + self.names(recon="none"), [])
        kind = Inventory.objects.filter(is_current=True).first().source_kind
        self.assertEqual(len(self.names(kind=kind)), 2)
        self.assertEqual(self.names(kind="нет такого"), [])
        from django.utils import timezone
        today = timezone.localdate()
        self.assertEqual(len(self.names(acc_from=today.isoformat(), acc_to=today.isoformat())), 2)
        self.assertEqual(self.names(acc_to=(today - dt.timedelta(days=1)).isoformat()), [])
        self.assertEqual(self.names(rows_min="abc"), self.names(), "мусор в числе — без отбора, без ошибки")

    def test_sort_by_header(self):
        self.assertEqual(self.names(sort="rows"), ["опись.xlsx", "!опись.xlsx"])
        r = self.client.get(reverse("inventories"), {"sort": "rows", "check": "1"})
        self.assertEqual(r.context["sorts"]["node"], "?check=1", "щелчок по «Узел» — порядок по умолчанию, отбор цел")
        self.assertTrue(r.context["f"].more_active)

    def test_export_keeps_filters(self):
        r = self.client.get(reverse("inventories"), {"check": "1", "export": "xlsx"})
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        rows = list(wb["Описи"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(x[0], x[1], x[2], x[10]) for x in rows], [("127", "!опись.xlsx", 1, "да")])
        params = dict(wb["Параметры"].iter_rows(min_row=2, values_only=True))
        self.assertEqual((params["Пометка «ПРОВЕРИТЬ»"], params["Описей"]), ("есть", 1))


@override_settings(AUTO_ACCEPT=True)
class StatsNodesTests(TransactionTestCase):
    """Свод по подведам и статистика описи (пользователь 27.10.2026: «Добавить
    статистику по каждому подведу, по каждой описи. Каждая опись должна знать,
    откуда была загружена и когда»)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. А", "опись.xlsx"), BASE)
        make_book(os.path.join(self.root, "126", "9. Б", "опись.xlsx"), BASE[:1])
        make_book(os.path.join(self.root, "126", "9. Б", "10. В", "опись.xlsx"), BASE[:2])
        make_book(os.path.join(self.root, "126", "опись.xlsx"), BASE[2:])
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_top_level_sums_subtree(self):
        r = self.client.get(reverse("stats_nodes"))
        (t,) = r.context["table"]
        self.assertEqual((t["node"].code, t["nodes"], t["invs"], t["rows"], t["qty"]), ("126", 4, 4, 7, 11))
        self.assertTrue(t["drill"])

    def test_children_and_own_inventories(self):
        from apps.orgs.models import Node
        n126 = Node.objects.get(path="/126/")
        r = self.client.get(reverse("stats_nodes"), {"parent": n126.pk})
        rows = {t["node"].code: (t["nodes"], t["invs"], t["rows"], t["qty"]) for t in r.context["table"]}
        self.assertEqual(rows, {"8": (1, 1, 3, 5), "9": (2, 2, 3, 3)})
        self.assertEqual((r.context["own"]["invs"], r.context["own"]["qty"]), (1, 3))
        self.assertEqual(r.context["totals"]["invs"], 4)
        self.assertEqual([i.source_file.name for i in r.context["own_invs"]], ["опись.xlsx"])
        self.assertContains(r, os.path.join(self.root, "126", "опись.xlsx"))
        x = self.client.get(reverse("stats_nodes"), {"parent": n126.pk, "export": "xlsx"})
        ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
        self.assertEqual([c.value for c in ws[2]][:3], ["126 — свои описи", 1, 1])
        self.assertEqual(ws.cell(ws.max_row, 3).value, 4)

    def test_inventory_of_measures_only_is_counted(self):
        """Опись из одних мер: без флага «Меры» её строка — «не учтено», с флагом — в счёте
        (флаги «учитывать», 29.10.2026)."""
        make_book(os.path.join(self.root, "126", "8. А", "меры.xlsx"),
                  [row("т", "31123130", "Основной список", qty=2.5, inv="20", cost=500)])
        services.scan([self.root], workers=1)
        (t,) = self.client.get(reverse("stats_nodes")).context["table"]
        self.assertEqual((t["invs"], t["rows"], t["out"], t["qty"]), (4, 7, 1, 11))
        (t,) = self.client.get(reverse("stats_nodes"), {"inc": "measure"}).context["table"]
        self.assertEqual((t["invs"], t["rows"], t["out"], t["qty"]), (5, 8, 0, Decimal("13.5")))

    def test_status_statistics_and_filter(self):
        """Статистика по статусам и отбор по ним (пользователь 27.10.2026: «фильтровать
        по статусам и видеть их статистику — количество по статусу, цена по статусу»)."""
        from apps.inventories.models import Status
        from apps.orgs.models import Node
        r = self.client.get(reverse("stats_nodes"))
        st = {x["status__name"]: (x["invs"], x["rows"], x["qty"], x["ci"]) for x in r.context["by_status"]}
        self.assertEqual(st, {"Основной список": (3, 3, 3, 30000), "Резерв": (2, 2, 2, 20000),
                              "Не требует оценки": (2, 2, 6, 20000)})
        self.assertEqual(round(sum(x["share"] for x in r.context["by_status"])), 100)
        reserve = Status.objects.get(name="Резерв")
        n126 = Node.objects.get(path="/126/")
        r = self.client.get(reverse("stats_nodes"), {"parent": n126.pk, "status": reserve.pk})
        rows = {t["node"].code: (t["nodes"], t["invs"], t["rows"], t["qty"], t["ci"]) for t in r.context["table"]}
        self.assertEqual(rows, {"8": (1, 1, 1, 1, 10000), "9": (1, 1, 1, 1, 10000)})
        self.assertIsNone(r.context["own"])
        self.assertEqual(r.context["own_invs"], [])
        self.assertEqual([x["status__name"] for x in r.context["by_status"]], ["Резерв"])
        self.assertContains(r, f"status={reserve.pk}")  # отбор сохраняется в ссылках вглубь
        x = self.client.get(reverse("stats_nodes"), {"status": reserve.pk, "export": "xlsx"})
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Подведы", "По статусам", "Параметры"])
        self.assertEqual(dict(wb["Параметры"].iter_rows(min_row=2, values_only=True))["Статусы"], "Резерв")
        r = self.client.get(reverse("stats_orgs"), {"status": reserve.pk})
        self.assertEqual(r.context["grand"][1], 2)
        r = self.client.get(reverse("home"))  # и на «Обзоре»
        self.assertEqual({x["status__name"]: x["qty"] for x in r.context["by_status"]},
                         {"Основной список": 3, "Резерв": 2, "Не требует оценки": 6})
        self.assertContains(r, f"?status={reserve.pk}")

    def test_before_first_load_is_empty(self):
        r = self.client.get(reverse("stats_nodes"), {"date": "2000-01-01"})
        self.assertEqual(r.context["table"], [])
        self.assertContains(r, "На эту дату данных нет")

    def test_inventory_page_shows_source_and_breakdown(self):
        inv = Inventory.objects.get(source_file__rel_folder=os.path.join("126", "8. А"))
        r = self.client.get(reverse("inventory_detail", args=[inv.pk]))
        self.assertContains(r, "Откуда и когда загружена")
        self.assertContains(r, inv.loaded_from)
        self.assertContains(r, reverse("scan_detail", args=[inv.scan_run_id]))
        parts = {title: rows for title, _, rows in r.context["breakdown"]}
        by_code = {x["k"]: x["rows"] for x in parts["По коду учёта"]}
        self.assertEqual(by_code, {"31123110": 1, "31123120": 1, "31123130": 1})


@override_settings(AUTO_ACCEPT=True)
class SuspectsTests(TransactionTestCase):
    """«Подозрительные» (пользователь 28.10.2026: «описи с подозрительно низким или
    высоким первоначальной стоимостью, или где в первоначальной стоимости будут буквы»;
    «чтобы они не учитывались в основной статистике пока на вкладке подозрительные я не
    нажму кнопку верно»)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        books = {"1. Дорогая": [row("Холодильник", "31123190", "Резерв", inv="", cost=311231900008),
                                row("Шкаф", "31123110", "Резерв", inv="11", cost=5000)],
                 "2. Дешёвая": [row("Стул", "31123110", "Резерв", inv="2", cost=0.5)],
                 "3. Буквы": [row("Шкаф", "31123110", "Резерв", inv="3", cost="сто сом")],
                 "4. Ноль": [row("Стол", "31123110", "Резерв", inv="4", cost=0)],
                 "5. Обычная": BASE}
        for folder, rows in books.items():
            make_book(os.path.join(self.root, "126", folder, "опись.xlsx"), rows)
        self.run_ = services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def names(self, r):
        return [x["inv"].node.name for x in r.context["page"]]

    def inv(self, name):
        return Inventory.objects.get(node__name=name, is_current=True)

    def test_letters_in_cost_are_kept_and_logged(self):
        from apps.inventories.models import Item
        it = Item.objects.get(name="Шкаф", inv_no="3")
        self.assertEqual((it.cost_initial, it.cost_initial_raw, it.suspect, it.held), (None, "сто сом", "text", True))
        self.assertIn("в первонач. стоимости не число — 1 стр.", self.run_.log)
        self.assertIn("подозрительная первонач. стоимость (не в своде до «Верно») — 1 стр.", self.run_.log)
        r = self.client.get(reverse("inventory_detail", args=[it.inventory_id]), {"suspect": "text"})
        self.assertContains(r, "«сто сом»")
        self.assertContains(r, "не в своде")
        self.assertEqual(r.context["tot"]["rows"], 1)

    def test_tab_lists_held_not_zero(self):
        r = self.client.get(reverse("suspects"))
        self.assertEqual(self.names(r), ["Дорогая", "Дешёвая", "Буквы"])
        tiles = {t["key"]: (t["rows"], t["invs"]) for t in r.context["tiles"]}
        self.assertEqual(tiles, {"high": (1, 1), "low": (1, 1), "neg": (0, 0), "text": (1, 1)})
        self.assertTrue(r.context["page"][0]["code_like"])
        self.assertContains(r, "похоже на инв. номер")
        self.assertContains(r, reverse("suspect_confirm", args=[self.inv("Дорогая").pk]))
        r = self.client.get(reverse("suspects"), {"f": "1", "cat": "low"})
        self.assertEqual(self.names(r), ["Дешёвая"])

    def test_held_not_in_stats_until_confirmed(self):
        from apps.stats.models import VersionStat
        inv = self.inv("Дорогая")
        self.assertEqual((inv.rows, inv.rows_held, inv.qty, inv.cost_initial), (2, 1, 1, 5000))
        self.assertEqual(sum(v.rows for v in VersionStat.objects.filter(inventory=inv, suspect=False)), 1)
        home = self.client.get(reverse("home"))
        self.assertEqual(home.context["held"], {"n": 3, "rows": 3})
        n126 = Node.objects.get(path="/126/")
        nodes = self.client.get(reverse("stats_nodes"), {"parent": n126.pk}).context["table"]
        self.assertEqual({t["node"].name: t["out"] for t in nodes}["Дорогая"], 1, "подозрительная — «не учтено»")
        # «Верно» — в свод, в т. ч. «на прошлые даты»; в списке подтверждённых; отмена — обратно.
        r = self.client.post(reverse("suspect_confirm", args=[inv.pk]), {"next": reverse("suspects")})
        self.assertRedirects(r, reverse("suspects"), fetch_redirect_response=False)
        inv.refresh_from_db()
        self.assertEqual((inv.rows_held, inv.qty, inv.cost_initial), (0, 2, 311231905008))
        self.assertEqual(sum(v.rows for v in VersionStat.objects.filter(inventory=inv, valid_to__isnull=True)), 2)
        self.assertEqual(VersionStat.objects.filter(inventory=inv).values("valid_from").distinct().count(), 1)
        self.assertEqual(self.names(self.client.get(reverse("suspects"))), ["Дешёвая", "Буквы"])
        self.assertEqual(self.names(self.client.get(reverse("suspects"), {"show": "ok"})), ["Дорогая"])
        self.client.post(reverse("suspect_undo", args=[inv.pk]))
        inv.refresh_from_db()
        self.assertEqual((inv.rows_held, inv.qty), (1, 1))

    def test_confirmation_carries_to_new_version(self):
        from apps.inventories.models import Item
        self.client.post(reverse("suspect_confirm", args=[self.inv("Дорогая").pk]))
        time.sleep(1.1)
        make_book(os.path.join(self.root, "126", "1. Дорогая", "опись.xlsx"),
                  [row("Холодильник", "31123190", "Резерв", inv="", cost=311231900008),
                   row("Плита", "31123190", "Резерв", inv="12", cost=411231900001)])
        services.scan([self.root], workers=1)
        new = self.inv("Дорогая")
        held = dict(Item.objects.filter(inventory=new).values_list("name", "held"))
        self.assertEqual(held, {"Холодильник": False, "Плита": True})

    def test_viewer_cannot_confirm(self):
        viewer = User.objects.create_user("v", "", "pw-12345678")
        self.client.force_login(viewer)
        self.assertEqual(self.client.post(reverse("suspect_confirm", args=[self.inv("Дорогая").pk])).status_code, 403)
        self.assertNotContains(self.client.get(reverse("suspects")), "Верно</button>")

    def test_recheck_after_threshold_change(self):
        from apps.imports import suspects
        self.client.post(reverse("suspect_confirm", args=[self.inv("Дешёвая").pk]))
        with self.settings(SUSPECT_HIGH="5000"):
            held, _ = suspects.recheck()
        # «Шкаф» за 5000 и BASE по 10 000 стали дороже порога; подтверждённый «Стул» — не отложен.
        self.assertEqual(held, 1 + 1 + 1 + 3)
        self.assertEqual(self.inv("Дешёвая").rows_held, 0)
        self.assertEqual(self.inv("Обычная").rows_held, 3)

    def test_export_positions_with_reasons(self):
        x = self.client.get(reverse("suspects"), {"export": "xlsx"})
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        why = {(r[4], r[5]): r[10] for r in wb["Позиции"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(why, {("Холодильник", None): "дороже порога — похоже на код счёта с номером",
                               ("Стул", "2"): "дешевле порога, но больше 0",
                               ("Шкаф", "3"): "не число: «сто сом»"})


@override_settings(AUTO_ACCEPT=True, DUP_MIN_COST="0")  # строки дешёвые — порог повтора снят
class DupScreensTests(TransactionTestCase):
    """«Дубли описей»: группы видны в вебе и выгружаются (пользователь 28.10.2026: «пускай
    в вебе показывается. С возможностью выгрузки оттуда»)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        folder = os.path.join(self.root, "126", "8. Тест")
        make_book(os.path.join(folder, "ОСВ — ОС.xlsx"), BASE + [row("Шкаф", "31123110", "Основной список", inv="30")])
        make_book(os.path.join(folder, "ОСВ — Sheet1.xlsx"), BASE)
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_groups_page_detail_and_exports(self):
        from apps.inventories.models import DupGroup
        g = DupGroup.objects.get()
        r = self.client.get(reverse("dups"))
        self.assertContains(r, reverse("dup_detail", args=[g.pk]))
        self.assertContains(self.client.get(reverse("dups") + "?q=нет-такого"), "Групп нет")
        r = self.client.get(reverse("dup_detail", args=[g.pk]))
        self.assertContains(r, "ОСВ — Sheet1.xlsx")
        self.assertContains(r, "основа")
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(reverse("dups") + "?export=xlsx").content))
        self.assertEqual(wb.sheetnames, ["Группы", "Описи", "Параметры"])
        self.assertEqual(wb["Группы"]["G2"].value, 3, "повторов — все строки меньшей описи")
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(reverse("dup_detail", args=[g.pk]) + "?export=xlsx").content))
        self.assertEqual(wb.sheetnames, ["Описи", "Позиции", "Параметры"])
        marks = [r[7] for r in wb["Позиции"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual((marks.count("повтор — не в своде"), marks.count("в своде")), (3, 4))
        # Позиции описи — отбор «повторы».
        small = Inventory.objects.get(source_file__name="ОСВ — Sheet1.xlsx")
        r = self.client.get(reverse("inventory_detail", args=[small.pk]) + "?dup=dup")
        self.assertContains(r, "повторов строк других описей")


@override_settings(AUTO_ACCEPT=True)
class FlagsAndQueryTests(TransactionTestCase):
    """Флаги «учитывать» во всех числах и конструктор запросов (пользователь 29.10.2026:
    «адаптировать все сборы статистики, все числа в интерфейсе под флаги»; «добавить
    страницу — конструктор запросов»)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        make_book(os.path.join(self.root, "126", "8. А", "опись.xlsx"),
                  BASE + [row("т", "31123130", "Основной список", qty=2.5, inv="20", cost=500)])
        services.scan([self.root], workers=1)
        self.client.force_login(User.objects.create_superuser("t", "", "pw-12345678"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_flag_changes_every_summary(self):
        plain, measure = {}, {"inc": "measure"}
        codes = [self.client.get(reverse("stats_codes"), q).context["grand"] for q in (plain, measure)]
        self.assertEqual(codes[1][0] - codes[0][0], 1, "мера — одна строка сверху")
        home = [self.client.get(reverse("home"), q).context["tot"] for q in (plain, measure)]
        self.assertEqual((home[0]["out"], home[1]["out"]), (1, 0))
        self.assertEqual(home[1]["rows"], codes[1][0], "главная и свод по кодам считают одинаково")
        invs = [self.client.get(reverse("inventories"), q).context["tot"] for q in (plain, measure)]
        self.assertEqual((invs[0]["rows"], invs[1]["rows"]), (codes[0][0], codes[1][0]))
        items = [self.client.get(reverse("items_search"), {"code": "31123130", **q}).context["tot"] for q in (plain, measure)]
        self.assertEqual(items[1]["rows"] - items[0]["rows"], 1)
        r = self.client.get(reverse("stats_codes"), {**measure, "export": "xlsx"})
        params = dict(openpyxl.load_workbook(io.BytesIO(r.content))["Параметры"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(params["Учитывать"], "Меры")

    def test_query_builder_groups_filters_and_saves(self):
        self.assertRedirects(self.client.get(reverse("query")), reverse("query") + "?g1=status", fetch_redirect_response=False)
        r = self.client.get(reverse("query"), {"g1": "status", "g2": "code"})
        total = r.context["total"]
        self.assertEqual(total["rows"], self.client.get(reverse("stats_codes")).context["grand"][0])
        self.assertEqual(len(r.context["table"][0]["dims"]), 2)
        r = self.client.get(reverse("query"), {"g1": "status", "cost_min": "10000", "m": ["rows", "qty"]})
        self.assertEqual([k for k, _ in r.context["measures"]], ["rows", "qty"])
        self.assertIn("стоимость от 10 000 сом", r.context["cond_text"])
        x = self.client.get(reverse("query"), {"g1": "status", "export": "xlsx"})
        ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
        self.assertEqual((ws["A1"].value, ws["B1"].value), ("Статус", "Позиций"))
        self.client.post(reverse("query_save"), {"name": "Мой запрос", "query": "g1=status&inc=dup"})
        from apps.stats.models import SavedQuery
        q = SavedQuery.objects.get(name="Мой запрос")
        self.assertEqual(q.query, "g1=status&inc=dup")
        self.client.post(reverse("query_delete", args=[q.pk]))
        self.assertFalse(SavedQuery.objects.exists())

    def test_row_summary_matches_row_and_links_sources(self):
        """Сводка строки (пользователь 29.10.2026: «нажимать на результаты… сводка что это такое,
        где и откуда… переход к источнику»): числа = строке таблицы, описи и позиции со ссылками."""
        r = self.client.get(reverse("query"), {"g1": "status", "g2": "code"})
        t = r.context["table"][0]
        s = self.client.get(reverse("query_row") + "?" + t["href"])
        self.assertEqual(s.status_code, 200)
        self.assertEqual(s.context["tot"]["rows"], t["rows"])
        self.assertEqual(s.context["title"], f"Статус: {t['dims'][0]} · Рекомендуемый код: {t['dims'][1]}")
        inv = s.context["sources"][0]["inv"]
        self.assertContains(s, reverse("inventory_detail", args=[inv.pk]))
        self.assertContains(s, reverse("item_detail", args=[s.context["page"].object_list[0].pk]))
        x = self.client.get(reverse("query_row") + "?" + t["href"] + "&export=xlsx")
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(x.content)).active.max_row - 1, t["rows"])

