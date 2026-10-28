"""Импорт описей: разбор, обход с хешами, приём, история «на дату».

Книги строятся здесь же — описи пользователя в репозиторий не кладутся.
Каждый сценарий ловит конкретную ошибку (в имени теста — какую).
"""

import datetime as dt
import os
import shutil
import tempfile
import time

import openpyxl
from openpyxl.styles import PatternFill
from django.test import TransactionTestCase, override_settings
from django.utils import timezone

from apps.inventories.models import Inventory, Item, Status
from apps.orgs.models import FolderMapping, Node
from apps.stats import services as stats

from . import dedup, services
from .models import ScanRun
from .reader import parse_book

HEAD = ["Основное средство", "Код учёта", "Количество", "Инв.Номер", "Первонач. Стоимость", "Накоп. Износ",
        "Балансовая стоимость", "Приоритет", "Материальное ответственное лицо", "Рекомендуемый код", "Ведомство",
        "Подведомство", "Модель средства", "Рабочее или нет", "Статус", "Причина", "Комментарий",
        "Тип объекта (определён)", "Полностью амортизирован", "Подразделение", "Код организации (42→31)",
        "Строка источника", "Сальдо нач.: первонач.", "Сальдо нач.: износ", "Сальдо нач.: балансовая",
        "Приход: первонач.", "Приход: износ", "Расход: первонач.", "Расход: износ", "Доп. сведения источника",
        "Отметка пользователя (02.10.2026)"]


def row(name, code, status, qty=1, inv="1", cost=10000, mark=""):
    return [name, 31122110, qty, inv, cost, 0, cost, 4, None, code, None, "Тестовое управление", None, None,
            status, "причина", "комментарий", "Офисное оборудование", "Нет", None, 31122110, 5, 0, 0, 0, 0, 0, 0, 0,
            "Папка источника: 126\\8. Тест; Файл / лист источника: ОСВ 18.11. 2025.xls / ПЕРЕЧЕНЬ", mark]


def make_book(path, rows, removed=(), fill_first=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Опись"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    if fill_first:
        ws["A2"].fill = PatternFill("solid", fgColor=fill_first)
    ws2 = wb.create_sheet("Вынесено")
    ws2.append(HEAD)
    for r in removed:
        ws2.append(r)
    sv = wb.create_sheet("Сводка")
    sv.append(["Тестовое управление"])
    sv.append(["Источник: !ПРОВЕРИТЬ — ОСВ 18.11. 2025 (1).xlsx — 1С ОСВ по ОС; лист «ПЕРЕЧЕНЬ»; иерархия: Счет / ОС"])
    sv.append(["Стоимости — сальдо на конец периода"])
    sv.append([])
    sv.append(["Сверка с итогами источника", "Итог источника: первонач.", "Сумма позиций: первонач.",
               "Итог источника: износ", "Сумма позиций: износ", "Совпадает"])
    sv.append(["Счёт 31122110", 100, 120, 0, 0, "НЕТ — расхождение в источнике"])
    pr = wb.create_sheet("Правила")
    pr.append(["Поле", "Как заполнено"])
    pr.append(["Код учёта", "Из источника"])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


BASE = [row("Компьютер", "31123130", "Основной список", inv="10"),
        row("Принтер", "31123120", "Резерв", qty=None, inv="М-205"),
        row("Стул", "31123110", "Не требует оценки", qty=3, inv="12", mark="красный")]


@override_settings(AUTO_ACCEPT=False)
class ReaderTests(TransactionTestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_all_fields_summary_rules_marks(self):
        """Все столбцы, «Сводка», «Правила», заливка и отметка; пустое количество — 1 с пометкой."""
        p = os.path.join(self.dir, "a.xlsx")
        make_book(p, BASE, removed=[row("Здание", "31112120", "Вынесено", inv="1")], fill_first="FFFF00")
        r = parse_book(p)
        self.assertEqual(r["error"], "")
        self.assertEqual(len(r["items"]), 4)
        first, printer, chair = r["items"][:3]
        self.assertEqual(first["name_fill"], "FFFF00")
        self.assertEqual(first["recommended_code"], "31123130")
        self.assertEqual(first["account_code"], "31122110")
        self.assertEqual(printer["qty"], "1")
        self.assertEqual(printer["qty_note"], "пусто — за 1")
        self.assertEqual(printer["inv_key"], "м205")
        self.assertEqual(chair["user_mark"], "красный")
        self.assertEqual(chair["user_mark_date"], "2026-10-02")
        self.assertIsNone(first["user_mark_date"], "дата отметки — только у строк с отметкой")
        self.assertEqual(r["source_name"], "ОСВ 18.11. 2025.xls", "имя исходного файла — из «Доп. сведений»")
        self.assertEqual(r["source_kind"], "1С ОСВ по ОС", "вид источника — после последнего « — »")
        self.assertEqual(r["source_date"], "2025-11-18")
        self.assertEqual(r["cost_basis"], "сальдо на конец периода")
        self.assertEqual(len(r["reconciliations"]), 1)
        self.assertEqual(r["rules"][0]["field"], "Код учёта")

    def test_money_in_quantity_counts_as_one(self):
        """Смотрим по количеству: сумма в «Количестве» (равна цене, больше 1000,
        переставленные столбцы) — 1 шт., исходное число — в пометке."""
        p = os.path.join(self.dir, "q.xlsx")
        swapped = row("Компьютер", "31123130", "Резерв", qty=19990, inv="5", cost=1)
        equal = row("Здание", "31112120", "Вынесено", qty=6824964, inv="6", cost=6824964)
        make_book(p, [swapped, equal, row("Кабель", "31122190", "Резерв", qty=12.5, inv="7")])
        items = parse_book(p)["items"]
        self.assertEqual([i["qty"] for i in items], ["1", "1", "12.50"])
        self.assertEqual([i["excluded"] for i in items], [False, False, True], "мера (12,5) — не в свод")
        self.assertIn("19990", items[0]["qty_note"])
        # столбцы переставлены (стоимость 1) — стоимость восстановлена из «Количества»
        self.assertEqual((items[0]["cost_initial"], items[0]["cost_balance"]), ("19990.00", "19990.00"))
        self.assertEqual(items[1]["cost_initial"], "6824964.00", "у равной цене стоимость не трогается")
        self.assertEqual(items[2]["qty_note"], "")

    def test_broken_book_is_an_error_not_a_crash(self):
        p = os.path.join(self.dir, "b.xlsx")
        open(p, "wb").write(b"not a zip")
        r = parse_book(p)
        self.assertIn("не открывается", r["error"])


@override_settings(AUTO_ACCEPT=False)
class ScanTests(TransactionTestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.folder = os.path.join(self.root, "126", "8. Тестовое ДНВТ")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def scan(self, **kw):
        return services.scan([self.root], workers=kw.get("workers", 1))

    def test_scan_skips_unchanged_and_copies(self):
        """Повторный обход не открывает книги; копия того же файла не разбирается."""
        make_book(os.path.join(self.folder, "опись.xlsx"), BASE)
        run = self.scan()
        self.assertEqual((run.found, run.parsed, run.pending, run.errors), (1, 1, 1, 0))
        run = self.scan()
        self.assertEqual((run.unchanged, run.parsed), (1, 0))
        shutil.copy(os.path.join(self.folder, "опись.xlsx"), os.path.join(self.root, "копия.xlsx"))
        run = self.scan()
        # Книга читается один раз (хеш из тех же байтов), копия опознаётся после разбора.
        self.assertEqual((run.moved, run.pending), (1, 0))
        self.assertIn("копия:", run.log)
        self.assertEqual(Inventory.objects.count(), 1)

    def test_inventory_knows_source_and_totals(self):
        """Опись помнит, откуда, каким обходом загружена и файл на тот момент;
        итоги — по позициям (меры — отдельно, в количество и деньги не входят)."""
        path = os.path.join(self.folder, "опись.xlsx")
        make_book(path, BASE + [row("т", "31123130", "Основной список", qty=2.5, inv="20", cost=500),
                                row("Шкаф", "", "Основной список", inv="21", cost=700)])
        run = self.scan()
        inv = Inventory.objects.get()
        st = os.stat(path)
        self.assertEqual((inv.loaded_via, inv.loaded_from, inv.scan_run_id, inv.loaded_by_id),
                         (Inventory.Via.FOLDER, path, run.pk, None))
        self.assertEqual(inv.loaded_where, path)
        self.assertEqual(inv.file_size, st.st_size)
        self.assertLess(abs(inv.file_mtime.timestamp() - st.st_mtime), 1)
        self.assertEqual((inv.rows, inv.rows_excluded, inv.rows_no_code, inv.codes), (5, 1, 1, 3))
        self.assertEqual((inv.qty, inv.cost_initial, inv.depreciation), (6, 30700, 0))
        # Файл переписан — у новой версии свой снимок, у прежней остался прежний.
        time.sleep(1.1)
        make_book(path, BASE)
        self.scan()
        old, new = Inventory.objects.order_by("pk")
        self.assertEqual(old.file_size, st.st_size)
        self.assertEqual(new.file_size, os.stat(path).st_size)
        self.assertNotEqual(old.scan_run_id, new.scan_run_id)

    def test_accept_creates_nodes_and_stats(self):
        make_book(os.path.join(self.folder, "опись.xlsx"), BASE)
        self.scan()
        inv = Inventory.objects.get(state=Inventory.State.PENDING)
        self.assertEqual(inv.node.path, "/126/8/", "узлы создаются при обходе — дерево видно до приёма")
        services.accept(inv)
        node = Node.objects.get(path="/126/8/")
        self.assertEqual((node.code, node.name, node.parent.code), ("8", "Тестовое ДНВТ", "126"))
        rows = {(r["code"], r["status__name"]): r["qty"] for r in stats.by_code(timezone.now())}
        self.assertEqual(rows[("31123110", "Не требует оценки")], 3)
        self.assertEqual(sum(rows.values()), 5)

    def test_history_on_date_uses_version_period(self):
        """Свод на дату — по версии, действовавшей тогда (граф APCS: debt-position-at-time-broken)."""
        p = os.path.join(self.folder, "опись.xlsx")
        make_book(p, BASE)
        self.scan()
        services.accept(Inventory.objects.get(state=Inventory.State.PENDING))
        between = timezone.now()
        time.sleep(1.1)
        make_book(p, BASE + [row("Сервер", "31123130", "Основной список", inv="99")])
        self.scan()
        new = Inventory.objects.get(state=Inventory.State.PENDING)
        services.accept(new)
        old = Inventory.objects.exclude(pk=new.pk).get(state=Inventory.State.ACCEPTED)
        self.assertFalse(old.is_current)
        self.assertIsNotNone(old.valid_to)
        total = lambda at: sum(r["qty"] for r in stats.by_code(at))
        self.assertEqual(total(between), 5)
        self.assertEqual(total(timezone.now()), 6)

    def test_resaved_book_with_same_rows_is_unchanged(self):
        """Пересохранённая книга (другой SHA-256, те же строки) новой версией не становится."""
        p = os.path.join(self.folder, "опись.xlsx")
        make_book(p, BASE)
        self.scan()
        services.accept(Inventory.objects.get(state=Inventory.State.PENDING))
        time.sleep(1.1)
        make_book(p, BASE)  # тот же состав, новая запись файла
        run = self.scan()
        self.assertEqual(run.same_content + run.moved, 1)
        self.assertFalse(Inventory.objects.filter(state=Inventory.State.PENDING).exists())

    def test_folder_without_code_needs_node_and_is_remembered(self):
        make_book(os.path.join(self.root, "Архив", "опись.xlsx"), BASE)
        self.scan()
        inv = Inventory.objects.get(state=Inventory.State.PENDING)
        self.assertFalse(inv.node_resolved)
        with self.assertRaises(ValueError):
            services.accept(inv)
        node = Node(code="200", name="Выбрано вручную")
        node.save()
        services.accept(inv, node=node)
        self.assertTrue(FolderMapping.objects.filter(folder="Архив", node=node).exists())

    def test_rejected_is_not_offered_again(self):
        make_book(os.path.join(self.folder, "опись.xlsx"), BASE)
        self.scan()
        inv = services.reject(Inventory.objects.get(state=Inventory.State.PENDING))
        self.assertEqual(Item.objects.filter(inventory=inv).count(), 0)
        os.utime(os.path.join(self.folder, "опись.xlsx"))  # тронули файл — содержимое то же
        run = self.scan()
        self.assertEqual(run.parsed, 0)

    def test_errors_go_to_log_file_per_book(self):
        """Ошибки — в журнал по файлам: путь книги и что не так (с листом и строкой)."""
        make_book(os.path.join(self.folder, "опись.xlsx"),
                  BASE + [row("Компьютер", "31123130", "Резерв", qty=19990, inv="5", cost=1)])
        os.makedirs(os.path.join(self.root, "127"))
        open(os.path.join(self.root, "127", "битая.xlsx"), "wb").write(b"x")
        run = self.scan()
        self.assertEqual((run.errors, run.warnings), (1, 1))
        text = open(run.log_file, encoding="utf-8-sig").read()
        self.assertIn("битая.xlsx: книга не открывается", text)
        self.assertIn("опись.xlsx: сумма в «Количестве» — 1 стр.", text)
        self.assertIn("Опись 5", text)
        os.remove(run.log_file)

    def test_long_cell_values_are_stored(self):
        """Длинные значения ячеек (код, инв. номер, модель) записываются — 6 книг падали
        на ограничении длины полей (DataError: value too long)."""
        long = row("Станок " + "x" * 600, "3" * 40, "Основной список", inv="И" * 300)
        long[12] = "М" * 700  # модель
        make_book(os.path.join(self.folder, "опись.xlsx"), BASE + [long])
        run = self.scan()
        self.assertEqual((run.errors, run.pending), (0, 1))
        self.assertTrue(Item.objects.filter(recommended_code="3" * 40, inv_no="И" * 300).exists())

    def test_failed_book_is_retried_without_file_change(self):
        """Книга, которая не записалась, при следующем обходе берётся заново, а не «без изменений»."""
        from apps.inventories.models import SourceFile
        make_book(os.path.join(self.folder, "опись.xlsx"), BASE)
        self.scan()
        Inventory.objects.all().delete()  # как после сбоя записи: файл известен, описи нет
        self.assertTrue(SourceFile.objects.exists())
        run = self.scan()
        self.assertEqual((run.unchanged, run.pending), (0, 1))

    def test_broken_file_counts_as_error(self):
        """Отказ не выглядит успехом (граф APCS: debt-import-swallows-failures)."""
        os.makedirs(self.folder)
        open(os.path.join(self.folder, "битая.xlsx"), "wb").write(b"x")
        run = self.scan()
        self.assertEqual(run.state, ScanRun.State.DONE_ERRORS)
        self.assertEqual(run.errors, 1)
        self.assertIn("битая.xlsx", run.log)

    def test_parallel_parse_in_processes(self):
        """Разбор в процессах (spawn на Windows) даёт те же описи."""
        for i in range(3):
            make_book(os.path.join(self.root, "126", f"{i + 1}. Подвед {i}", "опись.xlsx"),
                      BASE + [row(f"Вещь {i}", "31123130", "Резерв", inv=f"x{i}")])
        run = self.scan(workers=2)
        self.assertEqual((run.parsed, run.pending, run.errors), (3, 3, 0))
        self.assertEqual(Item.objects.count(), 12)


@override_settings(AUTO_ACCEPT=True)
class AutoAcceptTests(TransactionTestCase):
    """База не для ручной работы: обход сам принимает описи, узлы — из имён папок,
    служебная «обработанные» — не узел."""

    def setUp(self):
        self.root = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_scan_accepts_and_builds_nodes_from_folder_names(self):
        make_book(os.path.join(self.root, "106", "1. Достук", "обработанные", "опись.xlsx"), BASE)
        make_book(os.path.join(self.root, "1", "ГУ Дирекция", "опись.xlsx"), BASE + [row("Сервер", "31123130", "Резерв", inv="s1")])
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.accepted, run.pending), (2, 0))
        paths = sorted(Node.objects.filter(inventories__is_current=True).values_list("path", flat=True))
        self.assertEqual(paths[0], "/1/#" + str(Node.objects.get(name="ГУ Дирекция").pk) + "/")
        self.assertEqual(Node.objects.get(code="1", parent__code="106").name, "Достук")
        self.assertFalse(Node.objects.filter(name="обработанные").exists())
        self.assertEqual(sum(r["qty"] for r in stats.by_code(timezone.now())), 11)

    def test_measures_excluded_from_stats(self):
        """Материалы и запасы (меры) хранятся, но в свод штук не входят — «Исключаем»."""
        make_book(os.path.join(self.root, "106", "опись.xlsx"),
                  BASE + [row("Бензин", "31122190", "Резерв", qty=12.15, inv="b1", cost=600),
                          row("т", "31122190", "Резерв", qty=3, inv="t1", cost=20000)])
        services.scan([self.root], workers=1)
        self.assertEqual(sum(r["qty"] for r in stats.by_code(timezone.now())), 5)
        self.assertEqual(Item.objects.filter(excluded=True).count(), 2)


@override_settings(AUTO_ACCEPT=True)
class GoneFilesTests(TransactionTestCase):
    """Модуль нормализации при пересборке удаляет прежнюю книгу и меняет имя («!ПРОВЕРИТЬ — »,
    разбиение по листам): пропавшая книга снимается с актуальных — свод не двоит (21.10.2026)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.dir = os.path.join(self.root, "126", "8. Тест", "обработанные")
        make_book(os.path.join(self.dir, "опись.xlsx"), BASE)  # 5 шт.
        services.scan([self.root], workers=1)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def total(self, at=None):
        return sum(r["qty"] for r in stats.by_code(at or timezone.now()))

    def test_renamed_book_is_counted_once(self):
        before = timezone.now()
        os.remove(os.path.join(self.dir, "опись.xlsx"))
        make_book(os.path.join(self.dir, "!ПРОВЕРИТЬ — опись.xlsx"), BASE[:2])  # пересобрана: 2 шт.
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.missing, run.accepted), (1, 1))
        self.assertIn("снята с актуальных", run.log)
        cur = Inventory.objects.get(is_current=True)
        self.assertEqual(cur.source_file.name, "!ПРОВЕРИТЬ — опись.xlsx")
        self.assertEqual(self.total(), 2, "старая опись не считается рядом с новой")
        self.assertEqual(self.total(before), 5, "на прежнюю дату — как было")

    def test_moved_book_with_same_bytes_is_not_a_copy(self):
        """Перенос того же файла: двойник пропал — книга разбирается, а не пропускается как копия."""
        os.makedirs(os.path.join(self.root, "127", "обработанные"))
        os.rename(os.path.join(self.dir, "опись.xlsx"), os.path.join(self.root, "127", "обработанные", "опись.xlsx"))
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.moved, run.accepted), (0, 1))
        self.assertEqual(Inventory.objects.get(is_current=True).node.code, "127")
        self.assertEqual(self.total(), 5)

    def test_returned_book_is_current_again(self):
        keep = os.path.join(self.root, "опись.bak")
        os.rename(os.path.join(self.dir, "опись.xlsx"), keep)
        make_book(os.path.join(self.root, "127", "обработанные", "другая.xlsx"), BASE[:1])  # чтобы корень не был пуст
        services.scan([self.root], workers=1)
        self.assertEqual(self.total(), 1)
        os.rename(keep, os.path.join(self.dir, "опись.xlsx"))  # тот же файл, тот же размер и время
        services.scan([self.root], workers=1)
        self.assertEqual(self.total(), 6)

    def test_empty_or_unavailable_root_retires_nothing(self):
        os.remove(os.path.join(self.dir, "опись.xlsx"))  # в корне ни одной книги: сбой, а не «всё удалили»
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.missing, self.total()), (0, 5))
        self.assertIn("пропавшие с актуальных не снимаются", run.log)
        shutil.rmtree(self.root)
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.missing, self.total()), (0, 5))

    def test_unreadable_folder_is_not_gone(self):
        from apps.inventories.models import SourceFile
        sf = SourceFile.objects.get()
        notes = []
        n = services._retire_gone({sf.path: sf}, set(), [sf.root], [os.path.dirname(self.dir)], timezone.now(), notes)
        self.assertEqual((n, self.total()), (0, 5))


@override_settings(PROCESSED_FOLDERS=["обработанные"], AUTO_ACCEPT=True)
class ProcessedOnlyTests(TransactionTestCase):
    """Читаются только книги в «обработанные»; исходные рядом не открываются и в журнал не идут."""

    def setUp(self):
        self.root = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_only_processed_books_are_read(self):
        org = os.path.join(self.root, "115", "8. Тест")
        make_book(os.path.join(org, "обработанные", "опись.xlsx"), BASE)
        make_book(os.path.join(org, "исходная.xlsx"), BASE + [row("Лишнее", "31123130", "Резерв", inv="z")])
        open(os.path.join(org, "битая исходная.xlsx"), "wb").write(b"x")
        run = services.scan([self.root], workers=1)
        self.assertEqual((run.found, run.accepted, run.errors), (1, 1, 0))
        self.assertNotIn("исходн", run.log)
        self.assertEqual(Inventory.objects.get().node.path, "/115/8/")



@override_settings(AUTO_ACCEPT=True)
class RetireRootTests(TransactionTestCase):
    """Смена папки реестра: описи прежней папки снимаются с актуальных, ничего не удаляется."""

    def test_retire_root(self):
        from django.core.management import call_command
        old, new = tempfile.mkdtemp(), tempfile.mkdtemp()
        try:
            make_book(os.path.join(old, "126", "опись.xlsx"), BASE)
            make_book(os.path.join(new, "126", "опись.xlsx"), BASE[:2])
            services.scan([old], workers=1)
            services.scan([new], workers=1)
            # Обе описи одного узла, строки новой целиком в старой — повторы, в своде один раз
            # (дубли описей, 28.10.2026); раньше здесь был двойной счёт 7.
            self.assertEqual(sum(r["qty"] for r in stats.by_code(timezone.now())), 5, "до снятия — повторы в своде")
            call_command("retire_root", old, stdout=open(os.devnull, "w"))
            self.assertEqual(sum(r["qty"] for r in stats.by_code(timezone.now())), 2)
            self.assertEqual(Inventory.objects.count(), 2, "ничего не удалено")
        finally:
            shutil.rmtree(old, ignore_errors=True)
            shutil.rmtree(new, ignore_errors=True)


@override_settings(AUTO_ACCEPT=True)
class DedupTests(TransactionTestCase):
    """Дубли описей (apps.imports.dedup): лист той же книги и пересохранённый вариант у
    одного учреждения считались дважды (Департамент дорожного хозяйства, 28.10.2026)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.folder = os.path.join(self.root, "126", "8. Тестовое ДНВТ")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_similar_inventories_are_merged_as_multisets(self):
        stools = [row("Табурет", "31123110", "Основной список", inv="", cost=900) for _ in range(3)]
        full = BASE + stools + [row("Шкаф", "31123110", "Основной список", inv="30", cost=700)]
        # Вариант: табуретов два, у компьютера другая цена, принтер без номера.
        part = [row("Компьютер", "31123130", "Основной список", inv="10", cost=12000),
                row("принтер ", "31123120", "Резерв", qty=None, inv="", cost=10000),
                row("Стул", "31123110", "Не требует оценки", qty=3, inv="12"),
                row("Шкаф", "31123110", "Основной список", inv="30", cost=700)] + stools[:2]
        make_book(os.path.join(self.folder, "опись — ОС.xlsx"), full)
        make_book(os.path.join(self.folder, "опись — Sheet1.xlsx"), part)
        # Чужая опись другого учреждения с теми же строками — не дубль.
        make_book(os.path.join(self.root, "126", "9. Другое ДНВТ", "опись.xlsx"),
                  full + [row("Сейф", "31123110", "Основной список", inv="40")])
        services.scan([self.root], workers=1)
        big = Inventory.objects.get(source_file__name="опись — ОС.xlsx")
        small = Inventory.objects.get(source_file__name="опись — Sheet1.xlsx")
        other = Inventory.objects.get(source_file__name="опись.xlsx")
        self.assertTrue(big.is_current and small.is_current)
        big.refresh_from_db(), small.refresh_from_db(), other.refresh_from_db()
        self.assertIsNotNone(big.dup_group_id)
        self.assertEqual(big.dup_group_id, small.dup_group_id)
        self.assertIsNone(other.dup_group_id)
        # Повторы — в меньшей: принтер (по наименованию, у одной нет номера), стул и шкаф (по
        # номеру), два табурета; компьютер с другой ценой — своя строка, «сверить».
        dup = dict(Item.objects.filter(inventory=small, dup_of__isnull=False).values_list("name", "dup_step")
                   .order_by("name"))
        self.assertEqual(Item.objects.filter(inventory=small, dup_of__isnull=False).count(), 5)
        self.assertEqual(dup, {"принтер": 2, "Стул": 1, "Табурет": 2, "Шкаф": 1})
        self.assertEqual(Item.objects.get(inventory=small, name="Компьютер").dup_check, "cost")
        self.assertEqual((big.rows_dup, small.rows_dup), (0, 5))
        # В свод меньшей описи вошёл только компьютер.
        self.assertEqual(small.cost_initial, 12000)
        g = big.dup_group
        self.assertEqual((g.base_id, g.rows_before, g.rows_after), (big.pk, 13, 8))
        # Повторный проход по готовой базе даёт то же самое.
        dedup.run()
        small.refresh_from_db()
        self.assertEqual(small.rows_dup, 5)


@override_settings(AUTO_ACCEPT=True)
class DedupYearsTests(TransactionTestCase):
    """Годы описей в группе: опись старше самого нового года не в своде целиком
    («Кадамжай РайОО»: книги 2024 и 2025 годов, списанное к 2025-му оставалось в своде)."""

    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.folder = os.path.join(self.root, "126", "8. Тестовое ДНВТ")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_year_of(self):
        self.assertEqual(dedup.year_of(dt.date(2023, 5, 1), "опись 2025.xlsx"), 2023)
        self.assertEqual(dedup.year_of(None, "Инвентарь 2025 общ. — исфайрам — опись.xlsx"), 2025)
        self.assertEqual(dedup.year_of(None, "ОС 2023-2024 — опись.xlsx"), 2024)
        self.assertIsNone(dedup.year_of(None, "бпатп опись — 120250 — опись.xlsx"))

    def test_older_year_is_out_of_stats(self):
        # Дата источника читается из имени исходного файла в строках описи — у книги 2024 года свой.
        old_rows = [r[:29] + [r[29].replace("18.11. 2025", "18.11. 2024")] + r[30:]
                    for r in BASE + [row("Шкаф", "31123110", "Основной список", inv="30", cost=700)]]
        make_book(os.path.join(self.folder, "инвентаризация 2024 — опись.xlsx"), old_rows)
        make_book(os.path.join(self.folder, "Инвентарь 2025 — опись.xlsx"), BASE)
        services.scan([self.root], workers=1)
        old = Inventory.objects.get(source_file__name__contains="2024")
        new = Inventory.objects.get(source_file__name__contains="2025")
        self.assertEqual((old.year, new.year), (2024, 2025))
        self.assertEqual(old.outdated_by_id, new.pk)
        self.assertIsNone(new.outdated_by_id)
        # В своде — только опись 2025 года, шкаф (списан к 2025-му) не считается; повторов нет.
        self.assertEqual((old.cost_initial, new.cost_initial, new.rows_dup), (0, 30000, 0))
        self.assertEqual(sum(r["qty"] for r in stats.by_code(timezone.now())), 5)
        self.assertEqual(old.dup_group.base_id, new.pk)
