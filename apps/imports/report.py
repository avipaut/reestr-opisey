"""Отчёт обхода в Excel: ошибки с путями, ошибки данных, пропущенные файлы,
исключённые из свода позиции (пользователь 16.10.2026: «Сделай список ошибок
с путями и пропущенные в файл екселя»). Одна реализация — для команды
report_scan и кнопки «Выгрузить в Excel» на странице обхода."""

import io
import os

import openpyxl
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill

from apps.inventories.models import Item

from .models import ScanFile

BOLD, FILL = Font(bold=True), PatternFill("solid", fgColor="E8EDF3")


def _sheet(wb, title, head, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill, c.alignment = BOLD, FILL, Alignment(wrap_text=True, vertical="center")
    n = 0
    for r in rows:
        ws.append(r)
        n += 1
    if not n:
        ws.append(["— нет —"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return n


def _path(f):
    return os.path.join(f.root, f.rel_dir, f.name)


def build(run):
    """Книга отчёта по обходу → байты .xlsx."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    files = ScanFile.objects.filter(run=run).order_by("rel_dir", "name")
    S = ScanFile.Status
    n_err = _sheet(wb, "Ошибки", ["Путь к книге", "Причина"],
                   ([_path(f), f.note] for f in files.filter(status=S.ERROR).iterator()), [110, 80])
    n_data = _sheet(wb, "Ошибки данных", ["Путь к книге", "Итог файла", "Что не так (лист и строки)"],
                    ([_path(f), f.get_status_display(), f.note] for f in files.exclude(status=S.ERROR).exclude(note="")
                     .exclude(status=S.COPY).iterator()), [110, 18, 100])
    n_skip = _sheet(wb, "Пропущенные файлы", ["Путь к книге", "Почему пропущен", "Пояснение"],
                    ([_path(f), f.get_status_display(), f.note] for f in files.filter(
                        status__in=[S.IGNORED, S.COPY, S.UNCHANGED, S.SAME]).iterator()), [110, 36, 80])
    inv_ids = files.exclude(inventory_id=None).values_list("inventory_id", flat=True)
    excl = (Item.objects.filter(inventory_id__in=inv_ids, excluded=True).select_related("inventory__source_file")
            .order_by("inventory__source_file__path", "sheet", "row_no"))
    n_excl = _sheet(wb, "Исключённые позиции",
                    ["Путь к книге", "Лист", "Строка", "Инв. номер", "Основное средство", "Количество",
                     "Первонач. стоимость, сом", "Почему исключена"],
                    ([i.inventory.source_file.path, i.sheet, i.row_no, i.inv_no, i.name, float(i.qty),
                      float(i.cost_initial or 0), i.exclude_reason] for i in excl.iterator(chunk_size=2000)),
                    [100, 16, 9, 14, 45, 12, 16, 46])
    it = wb.create_sheet("Итог", 0)
    for r in [("Обход", run.pk), ("Папка", ", ".join(run.roots or [])),
              ("Начало", timezone.localtime(run.started_at).strftime("%d.%m.%Y %H:%M")),
              ("Конец", timezone.localtime(run.finished_at).strftime("%d.%m.%Y %H:%M") if run.finished_at else "идёт"),
              ("Итог", run.get_state_display()), ("Книг найдено", run.found), ("Считано", run.accepted),
              ("Ошибок (лист «Ошибки»)", n_err), ("Книг с ошибками данных (лист «Ошибки данных»)", n_data),
              ("Пропущено файлов (лист «Пропущенные файлы»)", n_skip),
              ("Позиций исключено из свода (лист «Исключённые позиции»)", n_excl),
              ("Пик памяти обхода, МБ", run.mem_peak_mb),
              ("Выгружено", timezone.localtime().strftime("%d.%m.%Y %H:%M"))]:
        it.append(list(r))
    it.column_dimensions["A"].width = 52
    it.column_dimensions["B"].width = 100
    for c in it["A"]:
        c.font = BOLD
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"ошибок": n_err, "ошибок данных": n_data, "пропущено": n_skip, "исключено": n_excl}
