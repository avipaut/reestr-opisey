"""Разбор нормализованной описи — без Django и базы.

Работает в процессах-читателях (spawn на Windows): модуль при импорте ничего
не делает, на вход — путь и байты, на выход — простые данные. Запись в базу —
в одном потоке и одной транзакцией (граф APCS: debt-threadpool-orm).
Формат — запись общего графа normalizovannaya-opis-format.
"""

import datetime as dt
import hashlib
import io
import re
from decimal import Decimal, InvalidOperation

import openpyxl

LIST_SHEETS = ("Опись", "Вынесено", "Ручная валидация")

# Поле позиции → возможные названия столбца (сравнение без регистра, «ё» = «е»).
COLUMNS = {
    "name": ["Основное средство", "Наименование"],
    "account_code": ["Код учёта", "Код учета"],
    "qty": ["Количество", "Кол-во"],
    "inv_no": ["Инв.Номер", "Инв. номер", "Инвентарный номер"],
    "cost_initial": ["Первонач. Стоимость", "Первоначальная стоимость"],
    "depreciation": ["Накоп. Износ", "Накопленный износ"],
    "cost_balance": ["Балансовая стоимость"],
    "priority": ["Приоритет"],
    "mol": ["Материальное ответственное лицо", "МОЛ"],
    "recommended_code": ["Рекомендуемый код"],
    "department_text": ["Ведомство"],
    "suborg_text": ["Подведомство"],
    "model": ["Модель средства", "Модель"],
    "working": ["Рабочее или нет"],
    "status_doc": ["Статус"],
    "reason": ["Причина"],
    "comment": ["Комментарий"],
    "object_type": ["Тип объекта (определён)", "Тип объекта"],
    "fully_depreciated": ["Полностью амортизирован"],
    "subdivision": ["Подразделение"],
    "org_code": ["Код организации (42→31)", "Код организации"],
    "source_row": ["Строка источника"],
    "saldo_begin_initial": ["Сальдо нач.: первонач."],
    "saldo_begin_dep": ["Сальдо нач.: износ"],
    "saldo_begin_balance": ["Сальдо нач.: балансовая"],
    "income_initial": ["Приход: первонач."],
    "income_dep": ["Приход: износ"],
    "expense_initial": ["Расход: первонач."],
    "expense_dep": ["Расход: износ"],
    "source_extra": ["Доп. сведения источника"],
}
REQUIRED = ("name", "status_doc", "recommended_code")
MONEY = {"cost_initial", "depreciation", "cost_balance", "saldo_begin_initial", "saldo_begin_dep",
         "saldo_begin_balance", "income_initial", "income_dep", "expense_initial", "expense_dep"}
CODES = {"account_code", "recommended_code", "org_code"}
USER_MARK = "отметка пользователя"  # «Отметка пользователя (02.10.2026)» — дата в шапке


def norm(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ").strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", s)


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def money(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    try:
        if isinstance(v, (int, float)):
            return str(round(Decimal(str(v)), 2))
        return str(round(Decimal(str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")), 2))
    except InvalidOperation:
        return None


def code(v) -> str:
    s = text(v)
    return s[:-2] if s.endswith(".0") else s


def inv_key(v) -> str:
    """Ключ позиции — сжатый инвентарный номер: без разделителей, латиница = кириллица
    (правило «Осмотра», запись графа opis-oborudovaniya-i-mehanizmov)."""
    s = norm(v)
    s = s.translate(str.maketrans("abcehkmoptxy", "авсенкмортху"))
    return re.sub(r"[^0-9a-zа-я]", "", s)


_DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.\s*(\d{4})")


def date_in(s):
    m = _DATE.search(s or "")
    if not m:
        return None
    try:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
    except ValueError:
        return None


def _header(ws, max_rows=20):
    """(номер строки шапки, {поле: столбец}, отметка: (столбец, дата), [неизвестные])."""
    syn = {f: {norm(x) for x in names} for f, names in COLUMNS.items()}
    for r_idx, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True), 1):
        cols, mark, unknown = {}, None, []
        for c_idx, v in enumerate(row):
            n = norm(v)
            if not n:
                continue
            hit = next((f for f, names in syn.items() if n in names and f not in cols), None)
            if hit:
                cols[hit] = c_idx
            elif n.startswith(USER_MARK):
                mark = (c_idx, date_in(text(v)))
            else:
                unknown.append((c_idx, text(v)))
        if all(f in cols for f in REQUIRED):
            return r_idx, cols, mark, unknown
    return None


def _fill(cell) -> str:
    try:
        f = cell.fill
        if f is not None and f.fill_type and f.fgColor is not None and f.fgColor.type == "rgb":
            rgb = f.fgColor.rgb or ""
            return "" if rgb in ("00000000", "FFFFFFFF") else rgb[-6:]
    except AttributeError:
        pass
    return ""


def _list_sheet(ws, out_items, unknown_cols):
    hdr = _header(ws)
    if not hdr:
        return False
    h_row, cols, mark, unknown = hdr
    for _, name in unknown:
        if name and name not in unknown_cols:
            unknown_cols.append(name)
    for r_idx, row in enumerate(ws.iter_rows(min_row=h_row + 1), h_row + 1):
        vals = [c.value for c in row]

        def get(f):
            i = cols.get(f)
            return vals[i] if i is not None and i < len(vals) else None

        if not any(text(get(f)) for f in ("name", "inv_no", "status_doc", "recommended_code")):
            continue
        it = {"sheet": ws.title, "row_no": r_idx}
        for f in COLUMNS:
            v = get(f)
            if f in MONEY:
                it[f] = money(v)
            elif f in CODES:
                it[f] = code(v)
            elif f in ("qty", "priority", "source_row", "fully_depreciated"):
                continue
            else:
                it[f] = text(v)
        q = money(get("qty"))
        it["qty"], it["qty_note"] = (q, "") if q is not None else ("1", "пусто — за 1" if not text(get("qty")) else "не число — за 1")
        p = get("priority")
        it["priority"] = int(p) if isinstance(p, (int, float)) or text(p).isdigit() else None
        sr = get("source_row")
        it["source_row"] = int(sr) if isinstance(sr, (int, float)) or text(sr).isdigit() else None
        fd = norm(get("fully_depreciated"))
        it["fully_depreciated"] = True if fd == "да" else False if fd == "нет" else None
        if mark:
            it["user_mark"] = text(vals[mark[0]]) if mark[0] < len(vals) else ""
            it["user_mark_date"] = mark[1] if it["user_mark"] else None
        else:
            it["user_mark"], it["user_mark_date"] = "", None
        it["name_fill"] = _fill(row[cols["name"]]) if cols["name"] < len(row) else ""
        it["extra"] = {name: text(vals[i]) for i, name in unknown if i < len(vals) and text(vals[i])}
        it["inv_key"] = inv_key(it["inv_no"])
        it["search"] = norm(" ".join(x for x in (it["name"], it["inv_no"], it["model"]) if x))
        h = hashlib.md5("\x1f".join(str(it.get(f) or "") for f in COLUMNS).encode("utf-8")).hexdigest()
        it["row_hash"] = h
        out_items.append(it)
    return True


def _summary(ws, info):
    lines = []
    recon, in_recon = [], False
    for row in ws.iter_rows(values_only=True):
        cells = [v for v in row if v not in (None, "")]
        if not cells:
            in_recon = False
            continue
        lines.append([text(v) if not isinstance(v, (int, float)) else v for v in cells])
        first = text(cells[0])
        if first.startswith("Сверка с итогами источника"):
            in_recon = True
            continue
        if in_recon and first.lower().startswith("счёт"):
            vals = list(row[1:6]) + [None] * 5
            recon.append({"account": first, "src_initial": money(vals[0]), "sum_initial": money(vals[1]),
                          "src_dep": money(vals[2]), "sum_dep": money(vals[3]), "matches": text(vals[4])})
            continue
        in_recon = False
        if first.startswith("Источник:"):
            s = first[len("Источник:"):].strip()
            parts = [p.strip() for p in s.split(";")]
            # «<файл> — <вид источника>»: в имени файла тоже бывает « — », делим по последнему.
            head = parts[0].rsplit(" — ", 1)
            info["summary_source"] = head[0].strip()
            if len(head) > 1:
                info["source_kind"] = head[1].strip()
            for p in parts[1:]:
                if p.startswith("лист"):
                    info["source_sheet"] = info["source_sheet"] or p.split("«", 1)[-1].rstrip("»").strip()
                elif p.startswith("иерархия"):
                    info["source_hierarchy"] = p.split(":", 1)[-1].strip()
        elif first.startswith("Стоимости"):
            info["cost_basis"] = first.split("—", 1)[-1].strip()
        elif len(lines) == 1:
            info["org_text"] = first
    info["summary_lines"] = lines
    info["reconciliations"] = recon


def _rules(ws, info):
    rules = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if row and text(row[0]):
            rules.append({"order": i, "field": text(row[0]), "text": text(row[1] if len(row) > 1 else "")})
    info["rules"] = rules


def parse_book(path: str, data: bytes = None) -> dict:
    """Разобрать книгу. Не бросает: ошибка — в поле error."""
    info = {"path": path, "error": "", "items": [], "unknown_columns": [], "org_text": "",
            "source_name": "", "source_sheet": "", "source_kind": "", "source_hierarchy": "",
            "source_folder": "", "cost_basis": "", "source_date": None, "summary_lines": [],
            "reconciliations": [], "rules": [], "sha256": "", "content_hash": ""}
    try:
        if data is None:
            with open(path, "rb") as f:
                data = f.read()
        info["sha256"] = hashlib.sha256(data).hexdigest()
        # read_only=False: нужны заливки ячеек (отметки пользователя цветом).
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # книга не открылась — это итог разбора, а не падение
        info["error"] = "книга не открывается: " + (str(e).splitlines() or [type(e).__name__])[0][:300]
        return info
    try:
        by_name = {norm(ws.title): ws for ws in wb.worksheets}
        found = False
        for name in LIST_SHEETS:
            ws = by_name.get(norm(name))
            if ws is not None:
                found = _list_sheet(ws, info["items"], info["unknown_columns"]) or found
        if not found:
            info["error"] = "не найдены листы описи с шапкой («Основное средство», «Статус», «Рекомендуемый код»)"
            return info
        if norm("Сводка") in by_name:
            _summary(by_name[norm("Сводка")], info)
        if norm("Правила") in by_name:
            _rules(by_name[norm("Правила")], info)
        extra = next((it["source_extra"] for it in info["items"] if it.get("source_extra")), "")
        m = re.search(r"Папка источника:\s*([^;]+)", extra)
        if m:
            info["source_folder"] = m.group(1).strip()
        m = re.search(r"Файл / лист источника:\s*(.+?)\s*/\s*([^/]+)$", extra)
        if m:
            info["source_name"] = info["source_name"] or m.group(1).strip()
            info["source_sheet"] = info["source_sheet"] or m.group(2).strip()
        info["source_name"] = info["source_name"] or info.pop("summary_source", "")
        info.pop("summary_source", None)
        info["source_date"] = date_in(info["source_name"]) or date_in(path.rsplit("\\", 1)[-1])
        info["content_hash"] = hashlib.sha256(
            "\n".join(sorted(it["row_hash"] for it in info["items"])).encode()).hexdigest()
        return info
    except Exception as e:
        info["error"] = f"ошибка разбора: {type(e).__name__}: {e}"[:500]
        return info
    finally:
        wb.close()
