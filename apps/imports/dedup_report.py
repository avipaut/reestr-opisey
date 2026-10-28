"""Выгрузка групп похожих описей в Excel — проверить объединение на глаз
(`manage.py dedup --report путь.xlsx`). Листы: «Группы», «Описи», «Сверить»;
шапка закреплена, автофильтр, перенос текста."""

from django.db.models import Count, Q, Sum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.inventories.models import DupGroup, Inventory, Item

MONEY = '#,##0.00'
INT = '#,##0'


def _sheet(wb, title, heads, rows, widths, formats=None):
    ws = wb.create_sheet(title)
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, f in (formats or {}).items():
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.number_format = f
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write(path):
    wb = Workbook()
    wb.remove(wb.active)
    groups = DupGroup.objects.select_related("node", "base__source_file").annotate(n=Count("inventories"))
    _sheet(wb, "Группы",
           ["Группа", "Узел", "Описей", "Основа (больше всего строк)", "Строк во всех описях", "Строк после объединения",
            "Повторов", "Первонач. стоимость во всех описях", "После объединения", "Разница"],
           [(g.pk, g.node.label, g.n,
             g.base.source_file.name if g.base else "", g.rows_before, g.rows_after, g.rows_before - g.rows_after,
             g.cost_before, g.cost_after, g.cost_before - g.cost_after)
            for g in groups.order_by("-cost_before")],
           [9, 40, 8, 50, 12, 12, 11, 20, 20, 20], {5: INT, 6: INT, 7: INT, 8: MONEY, 9: MONEY, 10: MONEY})

    invs = (Inventory.objects.filter(dup_group__isnull=False).select_related("node", "source_file", "dup_group")
            .annotate(s1=Count("items", filter=Q(items__dup_step=1)), s2=Count("items", filter=Q(items__dup_step=2)),
                      chk=Count("items", filter=~Q(items__dup_check="")),
                      dup_cost=Sum("items__cost_initial", filter=Q(items__dup_of__isnull=False)))
            .order_by("dup_group_id", "-rows"))
    _sheet(wb, "Описи",
           ["Группа", "Узел", "Файл", "Основа", "Строк", "Повторов: инв. номер", "Повторов: наименование",
            "Своих строк", "Сверить", "Стоимость повторов", "Путь"],
           [(i.dup_group_id, i.node.label if i.node else "", i.source_file.name, "да" if i.dup_group.base_id == i.pk else "",
             i.rows, i.s1, i.s2, i.rows - i.s1 - i.s2, i.chk, i.dup_cost or 0, i.loaded_where) for i in invs],
           [9, 34, 50, 8, 10, 12, 12, 11, 9, 18, 60], {5: INT, 6: INT, 7: INT, 8: INT, 9: INT, 10: MONEY})

    chk = (Item.objects.exclude(dup_check="").select_related("inventory__source_file", "inventory__node")
           .order_by("inventory__dup_group_id", "inv_key", "inventory_id")[:50000])
    _sheet(wb, "Сверить",
           ["Группа", "Узел", "Файл", "Строка", "Инв. номер", "Наименование", "Количество", "Первонач. стоимость",
            "Что отличается"],
           [(i.inventory.dup_group_id, i.inventory.node.label if i.inventory.node else "", i.inventory.source_file.name,
             i.source_row or i.row_no, i.inv_no, i.name, i.qty, i.cost_initial,
             "стоимость" if i.dup_check == "cost" else "количество") for i in chk],
           [9, 30, 45, 8, 16, 50, 11, 18, 14], {7: '#,##0.###', 8: MONEY})
    wb.save(path)
