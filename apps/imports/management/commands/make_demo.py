"""Тестовый набор для просмотра: описи в папках «учреждение / подвед» с
версиями в прошлом, прогон через настоящий импорт.

    manage.py make_demo [папка] [--orgs 8] [--seed 1]

Шаги: (1) книги первой версии → обход → приём на дату −120 дней;
(2) у части подведов — вторая версия (убраны, добавлены, сменили статус) →
приём на −45 дней; (3) у нескольких — третья версия и опись в папке без
кода — остаются «на просмотр». Данные выдуманы: названия, модели, цены.
Повторный запуск с той же папкой ничего не дублирует (хеши).
"""

import datetime as dt
import os
import random

import openpyxl
from openpyxl.styles import PatternFill
from django.conf import settings
from django.core.management.base import BaseCommand
from django.utils import timezone

from apps.imports import services
from apps.inventories.models import Inventory

HEAD = ["Основное средство", "Код учёта", "Количество", "Инв.Номер", "Первонач. Стоимость", "Накоп. Износ",
        "Балансовая стоимость", "Приоритет", "Материальное ответственное лицо", "Рекомендуемый код", "Ведомство",
        "Подведомство", "Модель средства", "Рабочее или нет", "Статус", "Причина", "Комментарий",
        "Тип объекта (определён)", "Полностью амортизирован", "Подразделение", "Код организации (42→31)",
        "Строка источника", "Сальдо нач.: первонач.", "Сальдо нач.: износ", "Сальдо нач.: балансовая",
        "Приход: первонач.", "Приход: износ", "Расход: первонач.", "Расход: износ", "Доп. сведения источника",
        "Отметка пользователя (02.10.2026)"]

# вид — (наименования, модели, рекомендуемый код, тип объекта, цена от-до)
KINDS = [
    (["Персональный компьютер", "Системный блок", "Ноутбук", "Монитор"], ["Dell OptiPlex 7090", "HP ProDesk 400", "Lenovo ThinkPad E14", "Samsung S24"],
     "31123130", "Компьютерное оборудование", (15000, 90000)),
    (["Принтер лазерный", "МФУ", "Копировальный аппарат", "Проектор"], ["Canon i-SENSYS MF445", "HP LaserJet M404", "Epson EB-X51", "Kyocera M2040"],
     "31123120", "Офисное оборудование", (8000, 60000)),
    (["IP-камера уличная", "Видеорегистратор", "Коммутатор"], ["Hikvision DS-2CD2047", "Dahua NVR4108", "TP-Link TL-SG1016"],
     "31123130", "Видеонаблюдение: IP-/веб-камеры, видеорегистраторы, контроль доступа", (3000, 40000)),
    (["Насос водяной", "Электродвигатель", "Генератор бензиновый", "Компрессор"], ["Grundfos CR 10", "АИР 100L4", "Huter DY6500", "Remeza СБ4"],
     "31122190", "Силовое / энергетическое оборудование", (20000, 400000)),
    (["Весы автомобильные", "Анализатор", "Измеритель давления"], ["ВА-15", "Mindray BC-30", "Testo 510"],
     "31122190", "Измерительный прибор / прибор учёта", (30000, 1800000)),
    (["Кондиционер", "Сейф", "Микрофон"], ["LG S12EQ", "Valberg ASM-63", "Shure SM58"],
     "31123190", "Механизмы и оборудование прочие", (6000, 70000)),
    (["Стол письменный", "Стул офисный", "Шкаф для документов"], ["", "", ""],
     "31123110", "Мебель", (1500, 25000)),
    (["Программа учёта (лицензия)", "Антивирус (лицензия)"], ["", ""],
     "31132190", "Нематериальный актив (ПО, лицензии)", (5000, 120000)),
]
REMOVED = [("Административное здание", "31112120", "Нежилое здание / помещение", (900000, 9000000)),
           ("Автомобиль легковой", "31121110", "Легковой автомобиль", (600000, 2500000))]
ORGS = ["Областная больница", "Департамент транспорта", "Управление образования", "Городской водоканал",
        "Центр стандартизации", "Районная администрация", "Учебный комбинат", "Дорожное управление",
        "Агентство связи", "Санитарная служба"]
SUBS = ["Центральный аппарат", "Северный филиал", "Южный филиал", "Склад", "Лаборатория", "Гараж", "Учебный центр"]
FILL = {"красный": "FFC7CE", "жёлтый": "FFEB9C", "оранжевый": "F4B084"}


def priority(cost):
    return 1 if cost >= 1e6 else 2 if cost >= 5e5 else 3 if cost >= 1e5 else 4 if cost >= 2e4 else 5


def status_of(kind_code, cost, rnd):
    if kind_code in ("31123110",):
        return "Не требует оценки", "Мебель — не требует оценки"
    if kind_code == "31132190":
        return "Не требует оценки", "Нематериальный актив: связь с оборудованием не подтверждена"
    if cost < 5000:
        return "Резерв", f"Стоимость единицы {cost:,.2f} ниже порога 5 000".replace(",", " ")
    if rnd.random() < 0.12:
        return "Резерв", "Вид имущества по наименованию не определён — код не установлен"
    return "Основной список", "Подлежит оценке"


def make_rows(rnd, n, org_name, start_inv):
    rows = []
    for i in range(n):
        names, models, code, otype, (lo, hi) = rnd.choice(KINDS)
        j = rnd.randrange(len(names))
        cost = round(rnd.uniform(lo, hi), 2)
        dep = round(cost * rnd.choice([0, 0.3, 0.6, 1.0]), 2)
        st, reason = status_of(code, cost, rnd)
        qty = rnd.choice([1] * 12 + [2, 3, 5])
        mark = rnd.choice(["", ""] * 20 + ["красный", "жёлтый", "оранжевый"])
        rows.append([f"{names[j]} {models[j % len(models)]}".strip(), 31122110, qty, str(start_inv + i), cost, dep,
                     round(cost - dep, 2), priority(cost), None, code, None, org_name, models[j % len(models)] or None,
                     None, st, reason, "", otype, "Да" if dep >= cost else "Нет", None, 31122110, 10 + i,
                     0, 0, 0, 0, 0, 0, 0, "Папка источника: демо; Файл / лист источника: ОСВ.xls / ПЕРЕЧЕНЬ", mark])
    return rows


def removed_rows(rnd, org_name, start_inv):
    out = []
    for k in range(rnd.randrange(0, 4)):
        name, code, otype, (lo, hi) = rnd.choice(REMOVED)
        cost = round(rnd.uniform(lo, hi), 2)
        out.append([name, int(code), 1, str(start_inv + k), cost, 0, cost, priority(cost), None, code, None,
                    org_name, None, None, "Вынесено", "Недвижимость / транспорт — вынесено из описи", "", otype,
                    "Нет", None, int(code), 5 + k, 0, 0, 0, 0, 0, 0, 0, "", ""])
    return out


def write_book(path, rows, removed, org_name, date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Опись"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
        if r[-1] in FILL:
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=FILL[r[-1]])
    ws2 = wb.create_sheet("Вынесено")
    ws2.append(HEAD)
    for r in removed:
        ws2.append(r)
    sv = wb.create_sheet("Сводка")
    sv.append([org_name])
    sv.append([f"Источник: ОСВ {date:%d.%m.%Y}.xls — 1С ОСВ по ОС; лист «ПЕРЕЧЕНЬ»; иерархия: Счет / Основное средство"])
    sv.append(["Стоимости — сальдо на конец периода"])
    pr = wb.create_sheet("Правила")
    pr.append(["Поле", "Как заполнено"])
    pr.append(["Статус", "Тестовый набор: по виду имущества и порогу 5 000"])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


class Command(BaseCommand):
    help = "Тестовый набор описей с историей версий — через настоящий импорт"

    def add_arguments(self, parser):
        parser.add_argument("folder", nargs="?", default=str(settings.BASE_DIR / ".dev" / "demo"))
        parser.add_argument("--orgs", type=int, default=8)
        parser.add_argument("--seed", type=int, default=1)

    def handle(self, folder, orgs, seed, **_):
        rnd = random.Random(seed)
        now = timezone.now()
        books = {}  # путь -> (строки, вынесено, имя)
        for o in range(orgs):
            org_code = 200 + o * 10
            for s in range(rnd.randrange(2, 6)):
                sub_name = SUBS[s % len(SUBS)]
                name = f"{ORGS[o % len(ORGS)]} — {sub_name}"
                path = os.path.join(folder, str(org_code), f"{s + 1}. {sub_name}", "опись.xlsx")
                if os.path.exists(path):
                    continue
                rows = make_rows(rnd, rnd.randrange(150, 900), name, 1000 * (s + 1))
                books[path] = (rows, removed_rows(rnd, name, 90000), name)
        if not books:
            self.stdout.write(f"В {folder} набор уже есть — ничего не добавлено")
            return
        # 1. первая версия: −120 дней
        for p, (rows, rem, name) in books.items():
            write_book(p, rows, rem, name, (now - dt.timedelta(days=125)).date())
        self._scan_accept(folder, now - dt.timedelta(days=120))
        # 2. вторая версия у ~40% подведов: −45 дней
        second = [p for p in books if rnd.random() < 0.4]
        for p in second:
            rows, rem, name = books[p]
            rows = [r[:] for r in rows if rnd.random() > 0.06]  # часть списали
            for r in rnd.sample(rows, min(len(rows), 12)):  # часть сменила статус
                r[14] = "Резерв" if r[14] == "Основной список" else "Основной список"
            rows += make_rows(rnd, rnd.randrange(5, 40), name, 50000)  # поступления
            books[p] = (rows, rem, name)
            write_book(p, rows, rem, name, (now - dt.timedelta(days=50)).date())
        self._touch(second)
        self._scan_accept(folder, now - dt.timedelta(days=45))
        # 3. ждут просмотра: третья версия у нескольких и опись в папке без кода
        pending = rnd.sample(list(books), min(3, len(books)))
        for p in pending:
            rows, rem, name = books[p]
            rows = [r[:] for r in rows[5:]] + make_rows(rnd, 8, name, 70000)
            write_book(p, rows, rem, name, now.date())
        self._touch(pending)
        write_book(os.path.join(folder, "Разное", "опись склада.xlsx"),
                   make_rows(rnd, 120, "Склад без привязки", 1), [], "Склад без привязки", now.date())
        run = services.scan([folder], workers=None)
        self.stdout.write(f"Готово: описей {len(books)}, вторых версий {len(second)}, "
                          f"на просмотр {run.pending}, ошибок {run.errors}. Папка: {folder}")

    def _touch(self, paths):
        t = timezone.now().timestamp()
        for p in paths:
            os.utime(p, (t, t))

    def _scan_accept(self, folder, at):
        run = services.scan([folder], workers=None)
        for inv in Inventory.objects.filter(state=Inventory.State.PENDING, source_file__root=os.path.normpath(folder)):
            services.accept(inv, at=at)
        self.stdout.write(f"  обход: разобрано {run.parsed}, принято на {timezone.localtime(at):%d.%m.%Y}")
