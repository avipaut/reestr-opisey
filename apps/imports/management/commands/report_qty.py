"""Отчёт для коллеги (модуль нормализации): книги и строки, где в «Количестве»
не штуки (пользователь 16.10.2026: «это баг коллеги видимо», «Файлы нужны.
пути», «Полный список такого пожалуйста в файл»).

    manage.py report_qty [файл.xlsx]

Листы: «По книгам» — полный путь и число строк каждого вида; «Исправлено» —
строки, посчитанные реестром за 1 шт. (исходное число — из пометки);
«Сомнительные» — не исправлены, решает человек: стоимость ≤ 1 при количестве
от 2 до 1000 (переставленные столбцы или настоящее количество) и дробное
количество. Только актуальные описи. Файлы описей не открываются — всё из базы.
"""

import os
import re
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand
from django.db.models import F
from django.db.models.functions import Floor
from openpyxl.styles import Alignment, Font, PatternFill

from apps.inventories.models import Item

KINDS = ["переставлены количество и стоимость", "количество равно стоимости", "количество = стоимость × 100",
         "количество больше 1000"]
DOUBT = ["стоимость ≤ 1, количество от 2 до 1000", "дробное количество"]


def _orig(note):
    m = re.search(r"в количестве (?:цена \()?([\d.]+)", note or "")
    return Decimal(m.group(1)) if m else None


def _kind(q, it):
    if it.cost_initial is not None and it.cost_initial <= 1:
        return KINDS[0]
    if q in (it.cost_initial, it.cost_balance):
        return KINDS[1]
    if it.cost_initial and q == it.cost_initial * 100:
        return KINDS[2]
    return KINDS[3]


def _row(it, qty, kind):
    sf = it.inventory.source_file
    return [sf.path, it.sheet, it.row_no, it.inv_no, it.name, float(qty), float(it.cost_initial or 0),
            float(it.cost_balance or 0), it.recommended_code, str(it.status or ""), kind]


class Command(BaseCommand):
    help = "Книги и строки с суммами в «Количестве» — для коллеги"

    def add_arguments(self, parser):
        parser.add_argument("out", nargs="?",
                            default=str(settings.BASE_DIR / ".dev" / "Количество-суммы — книги и строки.xlsx"))

    def handle(self, out, **_):
        cur = Item.objects.filter(inventory__is_current=True).select_related("inventory__source_file", "status")
        by_file = defaultdict(lambda: defaultdict(int))
        fixed, doubt = [], []
        for it in cur.exclude(qty_note="").order_by("inventory__source_file__path", "sheet", "row_no").iterator(
                chunk_size=2000):
            q = _orig(it.qty_note)
            k = _kind(q, it)
            fixed.append(_row(it, q, k))
            by_file[it.inventory.source_file.path][k] += 1
        plain = cur.filter(qty_note="")
        for k, qs in ((DOUBT[0], plain.filter(cost_initial__lte=1, qty__gt=1, qty__lte=1000)),
                      (DOUBT[1], plain.annotate(f=Floor("qty")).exclude(qty=F("f")))):
            for it in qs.order_by("inventory__source_file__path", "sheet", "row_no").iterator(chunk_size=2000):
                doubt.append(_row(it, it.qty, k))
                by_file[it.inventory.source_file.path][k] += 1

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "По книгам"
        ws.append(["Путь к книге", "Исправлено, строк", *KINDS, "Сомнительных, строк", *DOUBT])
        for path, d in sorted(by_file.items(), key=lambda x: -sum(x[1].values())):
            ws.append([path, sum(d[k] for k in KINDS), *[d[k] for k in KINDS],
                       sum(d[k] for k in DOUBT), *[d[k] for k in DOUBT]])
        head = ["Путь к книге", "Лист", "Строка", "Инв. номер", "Основное средство", "Количество в книге",
                "Первонач. стоимость, сом", "Балансовая стоимость, сом", "Рекомендуемый код", "Статус", "Вид"]
        for title, rows in (("Исправлено", fixed), ("Сомнительные", doubt)):
            w = wb.create_sheet(title)
            w.append(head)
            for r in rows:
                w.append(r)
        p = wb.create_sheet("Пояснение")
        for line in [
            ("Что это", "Строки нормализованных описей, где в столбце «Количество» не штуки."),
            ("Исправлено", "Реестр считает такую позицию за 1 шт.; «Количество в книге» — число как в файле."),
            ("Сомнительные", "Не исправлены: стоимость ≤ 1 при количестве 2–1000 (переставленные столбцы или "
                             "настоящее количество) и дробное количество (метры, килограммы?)."),
            ("Источник", "База реестра, актуальные описи; строка и лист — как в книге."),
        ]:
            p.append(list(line))
        bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8EDF3")
        for w in wb.worksheets:
            for c in w[1]:
                c.font, c.fill, c.alignment = bold, fill, Alignment(wrap_text=True, vertical="center")
            if w.title != "Пояснение":
                w.freeze_panes = "B2"
                w.auto_filter.ref = w.dimensions
            w.column_dimensions["A"].width = 100
        for w in (wb["Исправлено"], wb["Сомнительные"]):
            w.column_dimensions["E"].width = 45
            for col in "BCDFGHIJ":
                w.column_dimensions[col].width = 15
            w.column_dimensions["K"].width = 36
        for col in "BCDEFGHI":
            ws.column_dimensions[col].width = 16
        wb["Пояснение"].column_dimensions["B"].width = 100
        os.makedirs(os.path.dirname(out), exist_ok=True)
        wb.save(out)
        self.stdout.write(f"книг: {len(by_file)}; исправлено строк: {len(fixed)}; сомнительных: {len(doubt)}\n{out}")
